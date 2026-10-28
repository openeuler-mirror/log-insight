"""知识库文档解析: 按扩展名抽出纯文本 (一期不做 .doc)."""
from __future__ import annotations

import csv
import io
from pathlib import Path

SUPPORTED = {".txt", ".md", ".pdf", ".docx", ".xlsx", ".xls", ".csv"}


def parse_file(path: Path) -> str:
    ext = path.suffix.lower()
    if ext not in SUPPORTED:
        raise ValueError(f"不支持的格式: {ext}")
    if ext in {".txt", ".md"}:
        raw = path.read_bytes()
        for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
        return raw.decode("utf-8", errors="replace")
    if ext == ".pdf":
        return _pdf(path)
    if ext == ".docx":
        return _docx(path)
    if ext in {".xlsx", ".xls"}:
        return _xlsx(path)
    if ext == ".csv":
        return _csv(path)
    raise ValueError(f"不支持的格式: {ext}")


def _pdf(path: Path) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as e:
        raise ValueError("解析 PDF 需要安装 pymupdf") from e
    doc = fitz.open(path)
    try:
        return "\n".join(page.get_text() for page in doc)
    finally:
        doc.close()


def _docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise ValueError("解析 docx 需要安装 python-docx") from e
    d = Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(parts)


def _xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError as e:
        raise ValueError("解析 Excel 需要安装 openpyxl") from e
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"# {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(x.strip() for x in cells):
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)


def _csv(path: Path) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return "\n".join(" | ".join(row) for row in reader)
